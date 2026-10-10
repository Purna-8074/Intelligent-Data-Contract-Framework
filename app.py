from __future__ import annotations

import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st

st.set_page_config(
    page_title="Hospital Data Integration Studio",
    page_icon="🏥",
    layout="wide",
    initial_sidebar_state="expanded",
)

COMMON_COLUMNS = [
    "Hospital_Name", "Address", "City", "ZIP_Code", "Phone", "Website",
    "County_Status", "Longitude", "Latitude", "Location", "Data_Year",
    "Hospital_Type", "Bed_Count", "Quality_Score", "Record_Status",
]
NUMERIC_COLUMNS = ["Longitude", "Latitude", "Data_Year", "Bed_Count", "Quality_Score"]
TEXT_COLUMNS = [
    "Hospital_Name", "Address", "City", "Phone", "Website",
    "County_Status", "Hospital_Type", "Record_Status",
]

# These mappings reproduce the source notebook's field mappings.
SCHEMA_MAPPING = {
    "CSV": {
        "Hospital_Name": "Hospital_Name", "Address": "Address", "City": "City",
        "ZIP_Code": "ZIP_Code", "Phone": "Phone", "Website": "Website",
        "County_Status": "County_Status", "Longitude": "Longitude",
        "Latitude": "Latitude", "Location": "Location", "Data_Year": "Data_Year",
        "Hospital_Type": "Hospital_Type", "Bed_Count": "Bed_Count",
        "Quality_Score": "Inspection_Score", "Record_Status": None,
    },
    "JSON": {
        "Hospital_Name": "hospital_name", "Address": "street_address", "City": "city",
        "ZIP_Code": "postal_code", "Phone": "telephone", "Website": "web",
        "County_Status": "county_status", "Longitude": "lon", "Latitude": "lat",
        "Location": None, "Data_Year": "year", "Hospital_Type": "hospital_type",
        "Bed_Count": "beds", "Quality_Score": "rating", "Record_Status": None,
    },
    "XML": {
        "Hospital_Name": "HospitalName", "Address": "Street", "City": "Town",
        "ZIP_Code": "PostalCode", "Phone": "ContactPhone", "Website": "Website",
        "County_Status": "CountyStatus", "Longitude": "Longitude",
        "Latitude": "Latitude", "Location": None, "Data_Year": "Year",
        "Hospital_Type": "Type", "Bed_Count": "Beds", "Quality_Score": "QualityScore",
        "Record_Status": None,
    },
    "HTML": {
        "Hospital_Name": "Facility Name", "Address": "Street Address",
        "City": "Municipality", "ZIP_Code": "Postal", "Phone": "Contact",
        "Website": "Web Link", "County_Status": "County Flag", "Longitude": "Lon",
        "Latitude": "Lat", "Location": None, "Data_Year": "Reporting Year",
        "Hospital_Type": "Facility Category", "Bed_Count": "Beds Available",
        "Quality_Score": None, "Record_Status": "Record Status",
    },
    "KML": {
        "Hospital_Name": "name", "Address": None, "City": None, "ZIP_Code": None,
        "Phone": None, "Website": None, "County_Status": None, "Longitude": None,
        "Latitude": None, "Location": "coordinates", "Data_Year": None,
        "Hospital_Type": None, "Bed_Count": None, "Quality_Score": None,
        "Record_Status": None,
    },
    "KMZ": {
        "Hospital_Name": "name", "Address": None, "City": None, "ZIP_Code": None,
        "Phone": None, "Website": None, "County_Status": None, "Longitude": None,
        "Latitude": None, "Location": "coordinates", "Data_Year": None,
        "Hospital_Type": None, "Bed_Count": None, "Quality_Score": None,
        "Record_Status": None,
    },
}

ALIASES = {
    "Hospital_Name": ["hospital name", "facility name", "name", "hospital_name", "hospitalname"],
    "Address": ["address", "street", "street address", "street_address"],
    "City": ["city", "town", "municipality"],
    "ZIP_Code": ["zip", "zip code", "postal", "postal code", "postal_code", "zipcode"],
    "Phone": ["phone", "telephone", "contact", "contactphone", "contact phone"],
    "Website": ["website", "web", "web link", "url"],
    "County_Status": ["county status", "county_status", "county flag"],
    "Longitude": ["longitude", "lon", "lng"],
    "Latitude": ["latitude", "lat"],
    "Location": ["location", "coordinates"],
    "Data_Year": ["data year", "year", "reporting year"],
    "Hospital_Type": ["hospital type", "hospital_type", "type", "facility category"],
    "Bed_Count": ["bed count", "beds", "bed_count", "beds available"],
    "Quality_Score": ["quality score", "qualityscore", "rating", "inspection score", "inspection_score"],
    "Record_Status": ["record status", "record_status"],
}

def norm_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())

def detect_format(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    mapping = {".csv": "CSV", ".json": "JSON", ".xml": "XML", ".html": "HTML",
               ".htm": "HTML", ".kml": "KML", ".kmz": "KMZ"}
    if ext not in mapping:
        raise ValueError(f"Unsupported file extension: {ext or '(none)'}.")
    return mapping[ext]

def parse_kml_bytes(data: bytes) -> pd.DataFrame:
    root = ET.fromstring(data)
    rows = []
    # Namespace-agnostic parsing also supports KML documents with prefixes/default namespaces.
    for placemark in root.iter():
        if placemark.tag.split("}")[-1] != "Placemark":
            continue
        row = {}
        for child in placemark.iter():
            tag = child.tag.split("}")[-1]
            if tag == "name" and child.text and "name" not in row:
                row["name"] = child.text.strip()
            elif tag == "description" and child.text:
                row["description"] = child.text.strip()
            elif tag == "coordinates" and child.text:
                row["coordinates"] = child.text.strip()
        row.setdefault("coordinates", np.nan)
        rows.append(row)
    return pd.DataFrame(rows)

def parse_upload(uploaded_file, source_format: str) -> pd.DataFrame:
    data = uploaded_file.getvalue()
    name = uploaded_file.name
    if source_format == "CSV":
        return pd.read_csv(io.BytesIO(data))
    if source_format == "JSON":
        obj = json.loads(data.decode("utf-8-sig"))
        if isinstance(obj, dict):
            # Support a single record or common wrappers such as {"data": [...]}.
            candidate = next((v for v in obj.values() if isinstance(v, list)), None)
            obj = candidate if candidate is not None else [obj]
        return pd.json_normalize(obj)
    if source_format == "XML":
        root = ET.fromstring(data)
        records = []
        for element in root.iter():
            if element.tag.split("}")[-1].lower() in {"hospital", "record", "facility"}:
                record = {}
                for child in list(element):
                    record[child.tag.split("}")[-1]] = child.text
                if record:
                    records.append(record)
        if not records:
            # Fallback: parse repeated child elements from root.
            for element in list(root):
                record = {child.tag.split("}")[-1]: child.text for child in list(element)}
                if record:
                    records.append(record)
        return pd.DataFrame(records)
    if source_format == "HTML":
        tables = pd.read_html(io.BytesIO(data))
        if not tables:
            raise ValueError("No HTML table was found in this file.")
        return tables[0].copy()
    if source_format == "KML":
        return parse_kml_bytes(data)
    if source_format == "KMZ":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            kml_names = [n for n in archive.namelist()
                         if n.lower().endswith(".kml") and not n.startswith("__MACOSX/")]
            if not kml_names:
                raise ValueError("No .kml file was found inside this KMZ archive.")
            # Read the KML member in-memory; do not extract arbitrary archive paths.
            return parse_kml_bytes(archive.read(kml_names[0]))
    raise ValueError(f"Unsupported format: {source_format}")

def standardize_schema(df: pd.DataFrame, source_format: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    # Prefer the exact notebook mapping, then use normalized aliases for user-supplied variants.
    mapping = SCHEMA_MAPPING[source_format]
    lookup = {norm_name(c): c for c in df.columns}
    standardized = pd.DataFrame(index=df.index)
    mapping_rows = []
    for target in COMMON_COLUMNS:
        source = mapping.get(target)
        chosen = source if source is not None and source in df.columns else None
        if chosen is None:
            target_norm = norm_name(target)
            if target_norm in lookup:
                chosen = lookup[target_norm]
            else:
                for alias in ALIASES.get(target, []):
                    if norm_name(alias) in lookup:
                        chosen = lookup[norm_name(alias)]
                        break
        standardized[target] = df[chosen] if chosen is not None else np.nan
        mapping_rows.append({
            "Source_Format": source_format, "Common_Field": target,
            "Source_Field": chosen if chosen is not None else "Not available",
            "Mapping_Status": "Mapped" if chosen is not None else "Missing in source",
        })
    standardized["Source_Format"] = source_format
    standardized["Source_Record_ID"] = range(1, len(standardized) + 1)
    return standardized, pd.DataFrame(mapping_rows)

def make_missing_report(df: pd.DataFrame) -> pd.DataFrame:
    report = pd.DataFrame({
        "Column": df.columns,
        "Missing_Count": [int(df[c].isna().sum()) for c in df.columns],
    })
    report["Missing_Percentage"] = (
        report["Missing_Count"] / max(len(df), 1) * 100
    ).round(2)
    return report.sort_values("Missing_Count", ascending=False)

def preprocess_data(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict]:
    clean = raw.copy()
    before_rows = len(clean)
    before_missing = int(clean.isna().sum().sum())
    before_duplicates = int(clean.duplicated().sum())
    actions = []

    for col in TEXT_COLUMNS:
        clean[col] = clean[col].astype("string").str.strip()
    actions.append({"Step": "Text standardization", "Action": "Trimmed whitespace from text fields"})

    county_mapping = {
        "in county": "In-County", "in-county": "In-County",
        "out county": "Out-of-County", "out-of-county": "Out-of-County",
    }
    clean["County_Status"] = clean["County_Status"].apply(
        lambda x: county_mapping.get(str(x).strip().lower(), x) if pd.notna(x) else x
    )
    actions.append({"Step": "County category standardization", "Action": "Normalized county status labels"})

    type_mapping = {
        "general": "General Hospital", "emergency": "Emergency Center",
        "specialty": "Specialty Hospital",
    }
    clean["Hospital_Type"] = clean["Hospital_Type"].apply(
        lambda x: type_mapping.get(str(x).strip().lower(), x) if pd.notna(x) else x
    )
    actions.append({"Step": "Hospital type standardization", "Action": "Normalized common hospital type labels"})

    for col in NUMERIC_COLUMNS:
        clean[col] = pd.to_numeric(clean[col], errors="coerce")
    actions.append({"Step": "Numeric conversion", "Action": "Converted numeric fields; invalid text becomes missing"})

    invalid_beds = int((clean["Bed_Count"] < 0).sum())
    clean.loc[clean["Bed_Count"] < 0, "Bed_Count"] = np.nan
    invalid_quality = int(((clean["Quality_Score"] < 0) | (clean["Quality_Score"] > 100)).sum())
    clean.loc[(clean["Quality_Score"] < 0) | (clean["Quality_Score"] > 100), "Quality_Score"] = np.nan

    invalid_lon = (clean["Longitude"] < -180) | (clean["Longitude"] > 180)
    invalid_lat = (clean["Latitude"] < -90) | (clean["Latitude"] > 90)
    coordinate_errors = int(invalid_lon.sum() + invalid_lat.sum())
    clean.loc[invalid_lon, "Longitude"] = np.nan
    clean.loc[invalid_lat, "Latitude"] = np.nan

    clean["ZIP_Code"] = clean["ZIP_Code"].astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    zip_valid = clean["ZIP_Code"].str.match(r"^\d{5}$", na=False)
    invalid_zip = int((~zip_valid & clean["ZIP_Code"].notna()).sum())
    clean.loc[~zip_valid, "ZIP_Code"] = pd.NA

    def clean_phone(value):
        if pd.isna(value):
            return pd.NA
        digits = re.sub(r"\D", "", str(value))
        return digits if len(digits) == 10 else pd.NA
    clean["Phone"] = clean["Phone"].apply(clean_phone)
    clean["Website"] = clean["Website"].astype("string").str.strip().replace("", pd.NA)
    clean["Hospital_Name"] = clean["Hospital_Name"].astype("string").str.strip().str.replace(
        "National Institue", "National Institute", regex=False
    )

    rows_before_dedup = len(clean)
    clean = clean.drop_duplicates().reset_index(drop=True)
    duplicates_removed = rows_before_dedup - len(clean)

    imputation_rows = []
    for col in ["Bed_Count", "Quality_Score"]:
        missing_before = int(clean[col].isna().sum())
        median = clean[col].median()
        if pd.notna(median):
            clean[col] = clean[col].fillna(median)
        imputation_rows.append({
            "Column": col, "Missing_Before": missing_before,
            "Imputation_Method": "Median",
            "Imputation_Value": median if pd.notna(median) else "No median available",
            "Missing_After": int(clean[col].isna().sum()),
        })
    imputation_report = pd.DataFrame(imputation_rows)

    validation = []
    checks = [
        ("Negative Bed Count", int((clean["Bed_Count"] < 0).sum())),
        ("Quality Score 0-100", int(((clean["Quality_Score"] < 0) | (clean["Quality_Score"] > 100)).sum())),
        ("Latitude Range", int(((clean["Latitude"] < -90) | (clean["Latitude"] > 90)).sum())),
        ("Longitude Range", int(((clean["Longitude"] < -180) | (clean["Longitude"] > 180)).sum())),
        ("Duplicate Rows", int(clean.duplicated().sum())),
        ("Invalid ZIP format", int((clean["ZIP_Code"].notna() & ~clean["ZIP_Code"].str.match(r"^\d{5}$", na=False)).sum())),
        ("Invalid phone format", int((clean["Phone"].notna() & ~clean["Phone"].str.match(r"^\d{10}$", na=False)).sum())),
    ]
    for name, count in checks:
        validation.append({"Validation_Check": name, "Invalid_Count": count,
                           "Status": "PASS" if count == 0 else "FAIL"})
    validation_report = pd.DataFrame(validation)

    actions.extend([
        {"Step": "Invalid values", "Action": f"Invalid beds: {invalid_beds}; invalid quality scores: {invalid_quality}; invalid coordinate values: {coordinate_errors} set to missing"},
        {"Step": "ZIP validation", "Action": f"{invalid_zip} invalid ZIP values set to missing"},
        {"Step": "Phone standardization", "Action": "Kept only 10-digit phone numbers; others set to missing"},
        {"Step": "Hospital name standardization", "Action": "Applied notebook spelling correction for 'National Institue'"},
        {"Step": "Duplicate removal", "Action": f"Removed {duplicates_removed} exact duplicate rows"},
        {"Step": "Median imputation", "Action": "Filled missing Bed_Count and Quality_Score with column medians when available"},
    ])
    stats = {
        "rows_before": before_rows, "rows_after": len(clean),
        "columns": len(clean.columns), "missing_before": before_missing,
        "missing_after": int(clean.isna().sum().sum()),
        "duplicates_before": before_duplicates, "duplicates_removed": duplicates_removed,
        "invalid_beds": invalid_beds, "invalid_quality": invalid_quality,
        "invalid_coordinates": coordinate_errors, "invalid_zip": invalid_zip,
    }
    return clean, validation_report, imputation_report, {"stats": stats, "actions": pd.DataFrame(actions)}

def create_excel_report(raw: pd.DataFrame, clean: pd.DataFrame, source_summary: pd.DataFrame,
                        schema_report: pd.DataFrame, missing_report: pd.DataFrame,
                        validation_report: pd.DataFrame, imputation_report: pd.DataFrame,
                        actions: pd.DataFrame, before_after: pd.DataFrame,
                        dtype_report: pd.DataFrame, outlier_report: pd.DataFrame) -> bytes:
    final_display = clean.copy().fillna("NA")
    out = io.BytesIO()
    sheets = {
        "Raw_Combined_Data": raw, "Final_Clean_Data": final_display,
        "Source_Summary": source_summary, "Schema_Mapping": schema_report,
        "Data_Types": dtype_report, "Missing_Values": missing_report,
        "Outlier_Analysis": outlier_report, "Before_After": before_after,
        "Preprocessing_Log": actions, "Validation_Report": validation_report,
        "Imputation_Report": imputation_report,
    }
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for sheet, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet[:31], index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                from openpyxl.styles import PatternFill, Font, Alignment
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in ws.columns:
                max_len = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
                ws.column_dimensions[column_cells[0].column_letter].width = min(max_len + 2, 35)
            from openpyxl.styles import PatternFill
            for row in ws.iter_rows():
                for cell in row:
                    if str(cell.value).strip().upper() == "NA":
                        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    return out.getvalue()

st.markdown("""
<style>
.block-container {padding-top: 1.5rem; padding-bottom: 3rem; max-width: 1500px;}
.hero {padding: 1.3rem 1.5rem; border-radius: 18px; background: linear-gradient(120deg,#102a43,#176b87); color: white; margin-bottom: 1.2rem;}
.hero h1 {font-size: 2.1rem; margin: 0 0 .4rem 0;}
.hero p {margin: 0; opacity: .9;}
div[data-testid="stMetric"] {border: 1px solid rgba(128,128,128,.22); padding: 14px; border-radius: 14px;}
</style>
<div class="hero">
<h1>🏥 Hospital Data Integration Studio</h1>
<p>Upload heterogeneous hospital data, standardize schemas, explore data quality, preprocess records, validate, and export.</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.header("📂 Upload sources")
    uploads = st.file_uploader(
        "Choose one or more files",
        type=["csv", "json", "xml", "html", "htm", "kml", "kmz"],
        accept_multiple_files=True,
        help="Supported formats: CSV, JSON, XML, HTML, KML, KMZ.",
    )
    st.caption("The app processes uploaded files in memory; files are not saved to a server folder.")
    run_button = st.button("Process uploaded files", type="primary", use_container_width=True, disabled=not uploads)
    st.divider()
    st.markdown("**Processing flow**")
    st.caption("Upload → Parse → Schema map → Combine → EDA → Preprocess → Validate → Export")

if run_button and uploads:
    parsed, standardized, mapping_frames, errors = [], [], [], []
    progress = st.progress(0, text="Reading uploaded files…")
    for i, file in enumerate(uploads):
        try:
            fmt = detect_format(file.name)
            df = parse_upload(file, fmt)
            if df.empty:
                raise ValueError("The file was read, but no records were found.")
            std, mapping = standardize_schema(df, fmt)
            parsed.append({"filename": file.name, "format": fmt, "raw": df, "standardized": std})
            standardized.append(std)
            mapping_frames.append(mapping.assign(Source_File=file.name))
        except Exception as exc:
            errors.append({"File": file.name, "Error": str(exc)})
        progress.progress((i + 1) / len(uploads), text=f"Processed {i + 1} of {len(uploads)} file(s)")
    progress.empty()
    if errors:
        st.warning("Some files could not be processed. Review the details below; successfully read files will still be available.")
        st.dataframe(pd.DataFrame(errors), use_container_width=True, hide_index=True)
    if not standardized:
        st.error("No files were processed successfully. Check file formats and structure.")
        st.stop()

    raw = pd.concat(standardized, ignore_index=True)
    mapping_report = pd.concat(mapping_frames, ignore_index=True)
    source_summary = pd.DataFrame([
        {"File": p["filename"], "Source_Format": p["format"],
         "Input_Rows": len(p["raw"]), "Input_Columns": len(p["raw"].columns),
         "Mapped_Fields": int((mapping_report.loc[mapping_report["Source_File"] == p["filename"], "Mapping_Status"] == "Mapped").sum())}
        for p in parsed
    ])
    clean, validation_report, imputation_report, metadata = preprocess_data(raw)
    stats = metadata["stats"]
    missing_report = make_missing_report(raw)
    dtype_report = pd.DataFrame({"Column": raw.columns, "Data_Type": raw.dtypes.astype(str).values})
    outlier_rows = []
    for col in ["Bed_Count", "Quality_Score"]:
        vals = pd.to_numeric(raw[col], errors="coerce").dropna()
        if len(vals):
            q1, q3 = vals.quantile(.25), vals.quantile(.75)
            iqr = q3 - q1
            low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
            outlier_rows.append({"Variable": col, "Outlier_Count": int(((vals < low) | (vals > high)).sum()),
                                 "Lower_Bound": low, "Upper_Bound": high})
        else:
            outlier_rows.append({"Variable": col, "Outlier_Count": 0, "Lower_Bound": np.nan, "Upper_Bound": np.nan})
    outlier_report = pd.DataFrame(outlier_rows)
    before_after = pd.DataFrame([
        {"Metric": "Rows", "Before": len(raw), "After": len(clean)},
        {"Metric": "Missing cells", "Before": stats["missing_before"], "After": stats["missing_after"]},
        {"Metric": "Exact duplicate rows", "Before": stats["duplicates_before"], "After": int(clean.duplicated().sum())},
    ])
    st.session_state["hospital_result"] = {
        "parsed": parsed, "raw": raw, "clean": clean, "mapping": mapping_report,
        "source_summary": source_summary, "missing": missing_report, "dtypes": dtype_report,
        "validation": validation_report, "imputation": imputation_report,
        "actions": metadata["actions"], "before_after": before_after,
        "outliers": outlier_report, "errors": errors,
    }

result = st.session_state.get("hospital_result")
if not result:
    st.info("Upload your hospital files in the sidebar and select **Process uploaded files** to begin.")
    st.markdown("### Supported inputs")
    cols = st.columns(6)
    for col, fmt, detail in zip(cols, ["CSV", "JSON", "XML", "HTML", "KML", "KMZ"],
                                ["Tabular", "Records", "Elements", "HTML tables", "Map features", "Compressed KML"]):
        with col:
            st.markdown(f"#### `{fmt}`")
            st.caption(detail)
    st.markdown("### What this app does")
    st.markdown("- Uses the notebook's common 15-field schema and source-specific field mappings.\n"
                "- Combines the uploaded files and reports missingness, data types, duplicates, distributions, and outliers.\n"
                "- Applies notebook-aligned cleaning, range checks, ZIP/phone validation, duplicate removal, and median imputation.\n"
                "- Exports the raw standardized data and analysis reports to a formatted Excel workbook.")
    st.stop()

raw, clean = result["raw"], result["clean"]
st.caption(f"Current session: {len(result['parsed'])} file(s) processed · data is held in the app session.")
k1, k2, k3, k4, k5 = st.columns(5)
k1.metric("Uploaded files", len(result["parsed"]))
k2.metric("Combined records", f"{len(raw):,}")
k3.metric("Schema fields", len(COMMON_COLUMNS))
k4.metric("Missing cells before", f"{result['before_after'].loc[result['before_after']['Metric']=='Missing cells','Before'].iloc[0]:,}")
k5.metric("Validation checks passed", f"{int((result['validation']['Status']=='PASS').sum())}/{len(result['validation'])}")

tab_overview, tab_raw, tab_eda, tab_clean, tab_validation, tab_export = st.tabs([
    "📊 Overview", "🗂️ Integrated data", "🔎 EDA", "🧹 Preprocessing", "✅ Validation", "⬇️ Export"
])

with tab_overview:
    left, right = st.columns([1.05, 1])
    with left:
        st.subheader("Records by uploaded file")
        st.dataframe(result["source_summary"], use_container_width=True, hide_index=True)
        fig = px.bar(result["source_summary"], x="File", y="Input_Rows", color="Source_Format",
                     title="Input records per file", text_auto=True)
        st.plotly_chart(fig, use_container_width=True)
    with right:
        st.subheader("Processing pipeline")
        st.markdown("1. Parse each file using its format-specific reader.\n2. Map source columns to the shared hospital schema.\n3. Combine all standardized records.\n4. Profile data quality before cleaning.\n5. Apply preprocessing and validate the result.\n6. Download datasets and reports.")
        st.subheader("Before vs after")
        st.dataframe(result["before_after"], use_container_width=True, hide_index=True)
    st.caption("Counts are calculated from the files you uploaded in this session; they are not hardcoded sample results.")

with tab_raw:
    st.subheader("Integrated standardized dataset")
    st.caption("This table is the combined data after schema mapping and before preprocessing.")
    search = st.text_input("Filter rows by text", key="raw_filter")
    view = raw
    if search:
        mask = raw.astype("string").apply(lambda col: col.str.contains(search, case=False, na=False)).any(axis=1)
        view = raw.loc[mask]
    st.dataframe(view, use_container_width=True, hide_index=True, height=460)
    st.download_button("Download standardized CSV", raw.to_csv(index=False).encode("utf-8-sig"),
                       "hospital_standardized_raw.csv", "text/csv")

with tab_eda:
    st.subheader("Exploratory data analysis")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**Missing values by column**")
        miss = result["missing"].query("Missing_Count > 0").sort_values("Missing_Count")
        if miss.empty:
            st.success("No missing values were detected in the standardized data.")
        else:
            fig = px.bar(miss, x="Missing_Count", y="Column", orientation="h",
                         hover_data=["Missing_Percentage"], title="Missing values before preprocessing")
            st.plotly_chart(fig, use_container_width=True)
    with c2:
        st.markdown("**Records by source format**")
        by_format = raw.groupby("Source_Format", dropna=False).size().reset_index(name="Records")
        fig = px.pie(by_format, names="Source_Format", values="Records", hole=.45)
        st.plotly_chart(fig, use_container_width=True)
    st.markdown("**Missing-value report**")
    st.dataframe(result["missing"], use_container_width=True, hide_index=True)
    st.markdown("**Data types**")
    st.dataframe(result["dtypes"], use_container_width=True, hide_index=True)
    st.markdown("**Numeric summary**")
    numeric = raw.select_dtypes(include=np.number)
    if numeric.empty:
        st.info("No numeric fields were detected before preprocessing. Numeric conversion is included in the cleaning step.")
    else:
        st.dataframe(numeric.describe().T, use_container_width=True)
    st.markdown("**Outlier analysis (IQR rule)**")
    st.dataframe(result["outliers"], use_container_width=True, hide_index=True)
    if "Hospital_Type" in raw and raw["Hospital_Type"].notna().any():
        category_counts = raw["Hospital_Type"].fillna("Missing").value_counts().rename_axis("Hospital_Type").reset_index(name="Count")
        st.plotly_chart(px.bar(category_counts, x="Hospital_Type", y="Count", title="Hospital type distribution"), use_container_width=True)

with tab_clean:
    st.subheader("Preprocessed dataset")
    st.caption("Cleaning follows the uploaded notebook's main rules. Descriptive missing values are not invented.")
    st.dataframe(clean.fillna("NA"), use_container_width=True, hide_index=True, height=430)
    st.markdown("**Preprocessing log**")
    st.dataframe(result["actions"], use_container_width=True, hide_index=True)
    st.markdown("**Median imputation report**")
    st.dataframe(result["imputation"], use_container_width=True, hide_index=True)
    st.download_button("Download cleaned CSV", clean.fillna("NA").to_csv(index=False).encode("utf-8-sig"),
                       "hospital_final_clean_data.csv", "text/csv")

with tab_validation:
    st.subheader("Final data-quality validation")
    st.dataframe(result["validation"], use_container_width=True, hide_index=True)
    failures = result["validation"].query("Status == 'FAIL'")
    if failures.empty:
        st.success("All configured validation checks passed for this dataset.")
    else:
        st.warning(f"{len(failures)} validation check(s) still report invalid values. Review them before using the dataset.")
    st.caption("A PASS only means the specific check passed. Missing descriptive fields can remain in the final data by design.")

with tab_export:
    st.subheader("Export results")
    st.write("Create a formatted Excel workbook containing the standardized dataset, cleaned dataset, source summary, schema mapping, data types, missing-value report, outlier analysis, before/after comparison, preprocessing log, validation report, and imputation report.")
    excel_bytes = create_excel_report(
        raw, clean, result["source_summary"], result["mapping"], result["missing"],
        result["validation"], result["imputation"], result["actions"], result["before_after"],
        result["dtypes"], result["outliers"],
    )
    st.download_button(
        "📥 Download full Excel report",
        data=excel_bytes,
        file_name="Hospital_Data_Final_EDA_Preprocessing.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )
    st.download_button("Download schema mapping CSV", result["mapping"].to_csv(index=False).encode("utf-8-sig"),
                       "hospital_schema_mapping.csv", "text/csv")
    st.download_button("Download validation CSV", result["validation"].to_csv(index=False).encode("utf-8-sig"),
                       "hospital_validation_report.csv", "text/csv")
